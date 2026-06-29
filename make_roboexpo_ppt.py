from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# LG brand-inspired color palette
LG_RED    = RGBColor(0xCE, 0x0A, 0x2C)   # LG signature red
LG_NAVY   = RGBColor(0x1A, 0x2E, 0x4A)   # dark navy
LG_DARK   = RGBColor(0x1A, 0x1A, 0x1A)
LG_GRAY   = RGBColor(0x55, 0x55, 0x55)
LG_LGRAY  = RGBColor(0xF2, 0xF2, 0xF2)
LG_MGRAY  = RGBColor(0xDD, 0xDD, 0xDD)
LG_ACCENT = RGBColor(0xFF, 0x4C, 0x5C)   # lighter red
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LG_GOLD   = RGBColor(0xE8, 0xA8, 0x00)   # highlight gold

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

# ─── helpers ───────────────────────────────────────────────────────
def add_rect(slide, l, t, w, h, fill_rgb=None, line_rgb=None, line_width=None):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.line.fill.background()
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if line_rgb:
        shape.line.color.rgb = line_rgb
        if line_width:
            shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, l, t, w, h,
             font_size=18, bold=False, color=LG_DARK,
             align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(l, t, w, h)
    txb.word_wrap = wrap
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "맑은 고딕"
    return txb

def add_multiline(slide, lines, l, t, w, h,
                  font_size=16, bold=False, color=LG_DARK,
                  align=PP_ALIGN.LEFT, line_spacing=1.2):
    txb = slide.shapes.add_textbox(l, t, w, h)
    txb.word_wrap = True
    tf  = txb.text_frame
    tf.word_wrap = True
    first = True
    for line_text in lines:
        if first:
            p = tf.paragraphs[0]; first = False
        else:
            p = tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        run = p.add_run()
        run.text = line_text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "맑은 고딕"
    return txb

def slide_header(slide, title, subtitle=None):
    add_rect(slide, 0, 0, W, Inches(1.05), fill_rgb=LG_NAVY)
    add_rect(slide, 0, 0, Inches(0.08), Inches(1.05), fill_rgb=LG_RED)
    add_text(slide, title,
             Inches(0.25), Inches(0.1), Inches(10), Inches(0.65),
             font_size=26, bold=True, color=WHITE)
    if subtitle:
        add_text(slide, subtitle,
                 Inches(0.25), Inches(0.7), Inches(10), Inches(0.32),
                 font_size=13, color=LG_ACCENT)
    add_rect(slide, 0, Inches(7.2), W, Inches(0.3), fill_rgb=LG_NAVY)
    add_text(slide, "LG전자 생산기술원  |  2026 부산로봇엑스포 관람 보고서",
             Inches(0.2), Inches(7.22), Inches(9), Inches(0.25),
             font_size=10, color=WHITE)
    add_text(slide, "2026. 7. 3.",
             Inches(12.0), Inches(7.22), Inches(1.2), Inches(0.25),
             font_size=10, color=WHITE, align=PP_ALIGN.RIGHT)

def tag_box(slide, l, t, w, h, text, fill=LG_RED, txt_color=WHITE, fs=13):
    add_rect(slide, l, t, w, h, fill_rgb=fill)
    add_text(slide, text, l, t, w, h, font_size=fs, bold=True,
             color=txt_color, align=PP_ALIGN.CENTER)

def kpi_box(slide, l, t, w, h, number, label, num_color=LG_ACCENT):
    add_rect(slide, l, t, w, h, fill_rgb=LG_NAVY)
    add_rect(slide, l, t, Inches(0.06), h, fill_rgb=LG_RED)
    add_text(slide, number, l+Inches(0.15), t+Inches(0.08),
             w-Inches(0.2), Inches(0.58),
             font_size=28, bold=True, color=num_color, align=PP_ALIGN.CENTER)
    add_text(slide, label, l+Inches(0.1), t+Inches(0.65),
             w-Inches(0.15), Inches(0.35),
             font_size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════
# SLIDE 1 — 표지
# ═══════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill_rgb=LG_NAVY)
add_rect(sl, 0, 0, Inches(0.3), H, fill_rgb=LG_RED)
add_rect(sl, Inches(8.5), Inches(4.8), Inches(4.83), Inches(2.7), fill_rgb=LG_RED)

add_text(sl, "LG전자 생산기술원",
         Inches(0.6), Inches(0.9), Inches(10), Inches(0.5),
         font_size=18, color=LG_ACCENT)

add_text(sl, "전시회 관람 보고서",
         Inches(0.6), Inches(1.55), Inches(11), Inches(0.7),
         font_size=38, bold=True, color=WHITE)

add_text(sl, "2026 부산로봇기술산업전시회",
         Inches(0.6), Inches(2.35), Inches(11), Inches(0.55),
         font_size=22, bold=True, color=LG_ACCENT)

add_text(sl, "ROBO EXPO Busan 2026  ·  2026. 7. 1 – 7. 3  ·  부산 벡스코(BEXCO) 제2전시장 4홀",
         Inches(0.6), Inches(2.95), Inches(11), Inches(0.4),
         font_size=14, color=RGBColor(0xAA, 0xBB, 0xCC))

# 구분선
add_rect(sl, Inches(0.6), Inches(3.5), Inches(7.5), Inches(0.04), fill_rgb=LG_RED)

add_multiline(sl, [
    "주제: 로봇 인공지능 기반 기술과 혁신이 만드는 새로운 부산의 물결",
    "주관: 한국산업전람  |  동시개최: 부산오토매뉴팩 · 빅테크쇼",
], Inches(0.6), Inches(3.65), Inches(9), Inches(0.8),
   font_size=14, color=RGBColor(0xCC, 0xCC, 0xCC))

# KPIs
for i, (num, lbl) in enumerate([
    ("5개국", "참가국"),
    ("363부스", "전시 규모"),
    ("8,352㎡", "전시 면적"),
]):
    kpi_box(sl, Inches(0.6 + i*2.8), Inches(5.5), Inches(2.5), Inches(1.15), num, lbl)

add_text(sl, "작성: LG전자 생산기술원  |  관람일: 2026. 7. 1",
         Inches(0.6), Inches(6.85), Inches(8), Inches(0.35),
         font_size=12, color=RGBColor(0x88, 0x99, 0xAA))


# ═══════════════════════════════════════════════════
# SLIDE 2 — 전시회 개요
# ═══════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill_rgb=WHITE)
slide_header(sl, "전시회 개요", "2026 부산로봇엑스포 기본 현황")

# ── 왼쪽: 기본 정보 테이블 ──
add_rect(sl, Inches(0.3), Inches(1.2), Inches(5.8), Inches(5.7), fill_rgb=LG_LGRAY)
tag_box(sl, Inches(0.3), Inches(1.2), Inches(5.8), Inches(0.45),
        "행사 기본 정보", fill=LG_NAVY)

rows = [
    ("전시 기간", "2026. 7. 1(수) ~ 7. 3(금), 3일간"),
    ("전시 장소", "부산 벡스코(BEXCO) 제2전시장 4홀"),
    ("관람 시간", "7/1~7/2: 10:00~17:00  /  7/3: 10:00~16:00"),
    ("참가 규모", "5개국, 363개 부스 (8,352㎡)"),
    ("참가국", "한국, 독일, 일본, 중국, 말레이시아"),
    ("주관", "한국산업전람"),
    ("동시 개최", "부산오토매뉴팩 · 빅테크쇼\n2026 부산모빌리티쇼(BIMOS)"),
    ("전시 주제", "로봇 AI 기반 기술과 혁신이\n만드는 새로운 부산의 물결"),
]
for i, (label, val) in enumerate(rows):
    y = Inches(1.75 + i * 0.6)
    add_rect(sl, Inches(0.3), y, Inches(1.5), Inches(0.55), fill_rgb=LG_NAVY)
    add_text(sl, label, Inches(0.35), y+Inches(0.05), Inches(1.4), Inches(0.45),
             font_size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(sl, val, Inches(1.85), y+Inches(0.06), Inches(4.1), Inches(0.48),
             font_size=11, color=LG_DARK)

# ── 오른쪽: 전시 분야 구성 ──
add_rect(sl, Inches(6.4), Inches(1.2), Inches(6.6), Inches(5.7), fill_rgb=LG_LGRAY)
tag_box(sl, Inches(6.4), Inches(1.2), Inches(6.6), Inches(0.45),
        "전시 분야 구성", fill=LG_RED)

fields = [
    ("🤖 휴머노이드 로봇",   "이족·4족 보행 로봇, AI 자율행동"),
    ("🏭 산업용 로봇",       "다관절 로봇, 7축 용접로봇, SCARA"),
    ("🤝 협동로봇(Cobot)",  "중소기업 맞춤, 직접 교시 솔루션"),
    ("📦 물류 자동화",       "AMR, AGV, 대형 물류운반 드론"),
    ("🏗 스마트팩토리",      "디지털 트윈, AI 품질검사, MES 연동"),
    ("✈ 드론 / 무인기",     "산업용·수상·물류 드론"),
    ("🧩 로봇 부품·센서",   "감속기, 모터, 비전·힘 토크 센서"),
    ("🤖 서비스 로봇",       "서빙·조리 로봇, 매장 자동화 솔루션"),
]
for i, (field, desc) in enumerate(fields):
    y = Inches(1.75 + i * 0.6)
    add_rect(sl, Inches(6.4), y, Inches(2.4), Inches(0.52), fill_rgb=LG_NAVY)
    add_text(sl, field, Inches(6.45), y+Inches(0.06), Inches(2.3), Inches(0.42),
             font_size=11, bold=True, color=WHITE)
    add_text(sl, desc, Inches(8.85), y+Inches(0.06), Inches(4.1), Inches(0.48),
             font_size=11, color=LG_DARK)


# ═══════════════════════════════════════════════════
# SLIDE 3 — 주요 참가기업 & 전시 하이라이트
# ═══════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill_rgb=WHITE)
slide_header(sl, "주요 참가기업 및 전시 하이라이트",
             "5개국 363부스 — 글로벌 탑티어 기업 총출동")

companies = [
    {
        "name": "유니트리 로보틱스\n(Unitree Robotics)",
        "country": "중국",
        "highlight": "최대 주목 부스",
        "items": [
            "H1·G1 휴머노이드 실물 시연",
            "4족 보행 Go2 산업 적용 사례",
            "AI 자율보행 · 장애물 회피 시연",
            "가격 경쟁력 압도적 — 국내 대비",
        ],
        "tag_color": LG_RED,
    },
    {
        "name": "DJI",
        "country": "중국",
        "highlight": "드론·무인기 대표",
        "items": [
            "산업용 드론 · 대형 물류운반 드론",
            "공장 내 부품 이송 운용 시나리오",
            "수상 드론(해양 모니터링·물류)",
            "3D 공간 물류 동선 최적화 제안",
        ],
        "tag_color": LG_NAVY,
    },
    {
        "name": "협동로봇\n국내 전문업체군",
        "country": "한국",
        "highlight": "Cobot 솔루션",
        "items": [
            "직접 교시(Direct Teaching) 보편화",
            "다품종 소량 대응 유연 셀 구성",
            "힘·토크 센서 기반 정밀 협업",
            "중소 제조현장 맞춤형 패키지",
        ],
        "tag_color": RGBColor(0x1A, 0x7A, 0x3C),
    },
    {
        "name": "AMR/AGV\n자율물류 업체군",
        "country": "한국·독일·일본",
        "highlight": "물류 자동화",
        "items": [
            "SLAM 기반 자율 맵핑 실용화",
            "다중 로봇 Fleet Management SW",
            "MES 연동 실시간 동선 재설정",
            "좁은 통로·인원 혼재 안전 알고리즘",
        ],
        "tag_color": RGBColor(0x7B, 0x3F, 0x9A),
    },
]

for i, c in enumerate(companies):
    x = Inches(0.25 + i * 3.25)
    add_rect(sl, x, Inches(1.2), Inches(3.1), Inches(5.7), fill_rgb=LG_LGRAY)
    # 헤더
    add_rect(sl, x, Inches(1.2), Inches(3.1), Inches(0.5), fill_rgb=c["tag_color"])
    add_text(sl, c["country"], x+Inches(0.08), Inches(1.22), Inches(0.7), Inches(0.46),
             font_size=10, bold=True, color=WHITE)
    add_text(sl, c["highlight"], x+Inches(1.5), Inches(1.22), Inches(1.55), Inches(0.46),
             font_size=10, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    # 기업명
    add_text(sl, c["name"], x+Inches(0.1), Inches(1.75), Inches(2.9), Inches(0.65),
             font_size=16, bold=True, color=LG_NAVY)
    # 구분선
    add_rect(sl, x+Inches(0.1), Inches(2.43), Inches(2.9), Inches(0.03), fill_rgb=c["tag_color"])
    # 내용
    add_multiline(sl, c["items"], x+Inches(0.12), Inches(2.5), Inches(2.88), Inches(3.8),
                  font_size=12, color=LG_DARK)

# 하단 요약
add_rect(sl, Inches(0.25), Inches(6.95), Inches(12.8), Inches(0.32), fill_rgb=LG_RED)
add_text(sl,
         "핵심 관찰: 중국 기업(유니트리·DJI)의 기술 수준이 글로벌 선두권 진입 — 가격 경쟁력과 기술 완성도 동시 충격",
         Inches(0.4), Inches(6.97), Inches(12.5), Inches(0.28),
         font_size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════
# SLIDE 4 — LG생산기술원 관점 (1/2): 협동로봇·AMR
# ═══════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill_rgb=WHITE)
slide_header(sl, "LG생산기술원 관점 주요 발견 ①",
             "협동로봇(Cobot)  &  AMR 물류 자동화")

# ── 왼쪽: 협동로봇 ──
add_rect(sl, Inches(0.25), Inches(1.15), Inches(6.3), Inches(5.85), fill_rgb=LG_LGRAY)
tag_box(sl, Inches(0.25), Inches(1.15), Inches(6.3), Inches(0.5),
        "협동로봇(Cobot) — 제조 현장 적용 확대", fill=LG_RED)

add_text(sl, "관찰 내용",
         Inches(0.4), Inches(1.75), Inches(6.0), Inches(0.35),
         font_size=13, bold=True, color=LG_NAVY)
add_multiline(sl, [
    "• 직접 교시(Direct Teaching) UI 보편화 — 재프로그래밍 없이 현장 전환",
    "• 힘·토크 센서 내장으로 복잡한 케이블/부품 조립 대응 가능",
    "• 다품종 소량 환경에서 유연 셀 구성 능력 크게 향상",
    "• 인간-로봇 공존 안전 규격(ISO/TS 15066) 기반 제품 표준화",
], Inches(0.4), Inches(2.12), Inches(6.0), Inches(1.6),
   font_size=12, color=LG_DARK)

add_rect(sl, Inches(0.4), Inches(3.78), Inches(6.0), Inches(0.04), fill_rgb=LG_RED)
add_text(sl, "LG생산기술원 시사점",
         Inches(0.4), Inches(3.85), Inches(6.0), Inches(0.35),
         font_size=13, bold=True, color=LG_RED)
add_multiline(sl, [
    "▶ 가전·전장 조립공정 Human-Robot Collaboration 확대 검토",
    "▶ 힘·토크 센서 기반 Cobot의 LG 공정 적합성 평가 착수",
    "▶ 생산기술원 공정 최적화 노하우를 Cobot 티칭에 내재화",
    "▶ 단기 파일럿 대상 공정 선정 → PoC 기획 (6개월 내)",
], Inches(0.4), Inches(4.22), Inches(6.0), Inches(2.2),
   font_size=12, color=LG_DARK)

# ── 오른쪽: AMR ──
add_rect(sl, Inches(6.8), Inches(1.15), Inches(6.3), Inches(5.85), fill_rgb=LG_LGRAY)
tag_box(sl, Inches(6.8), Inches(1.15), Inches(6.3), Inches(0.5),
        "AMR 기반 공장 내 물류 자동화", fill=LG_NAVY)

add_text(sl, "관찰 내용",
         Inches(6.95), Inches(1.75), Inches(6.0), Inches(0.35),
         font_size=13, bold=True, color=LG_NAVY)
add_multiline(sl, [
    "• SLAM 기반 자율 맵핑 — 고정 레일 없이 공장 환경 동적 인식",
    "• 다중 AMR Fleet Management SW 완성도 대폭 향상",
    "• MES 연동 → 생산 계획 변경 시 실시간 동선 재설정",
    "• 충전 자율 귀환, 배터리 최적 관리, 인원 혼재 안전 알고리즘",
], Inches(6.95), Inches(2.12), Inches(6.0), Inches(1.6),
   font_size=12, color=LG_DARK)

add_rect(sl, Inches(6.95), Inches(3.78), Inches(6.0), Inches(0.04), fill_rgb=LG_RED)
add_text(sl, "LG생산기술원 시사점",
         Inches(6.95), Inches(3.85), Inches(6.0), Inches(0.35),
         font_size=13, bold=True, color=LG_RED)
add_multiline(sl, [
    "▶ 창원·구미 공장 AGV → AMR 전환 가속화 시점 재검토",
    "▶ 계열사 로보스타 AGV/AMR 솔루션과의 연계 강화 방안 수립",
    "▶ LG생산기술원 자체 개발 MM(웨이퍼 이송 로봇) 기술의\n    공장 물류 영역 확장 가능성 검토",
    "▶ AMR 도입 대상 공장 선정 및 타당성 분석 착수",
], Inches(6.95), Inches(4.22), Inches(6.0), Inches(2.2),
   font_size=12, color=LG_DARK)


# ═══════════════════════════════════════════════════
# SLIDE 5 — LG생산기술원 관점 (2/2): 휴머노이드·AI팩토리·7축 용접
# ═══════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill_rgb=WHITE)
slide_header(sl, "LG생산기술원 관점 주요 발견 ②",
             "휴머노이드 로봇  ·  AI 스마트팩토리  ·  7축 용접로봇")

cards = [
    {
        "title": "🤖 휴머노이드 로봇",
        "subtitle": "미래 제조 환경의 판도 변화",
        "obs": [
            "유니트리 G1 실시간 시연: 부품 집기·이송·조립",
            "비정형 작업 대응력 및 공간 활용성 우수",
            "작업 속도·정확도는 전용 로봇 대비 아직 부족",
            "중국산 상용 제품 가격 급속 하락 추세",
        ],
        "tips": [
            "단기(1~3년): 물류·자재 공급 등 비정밀 작업 파일럿",
            "중기(3~5년): 조립 보조·검사 등 반정밀 작업 확대",
            "장기: 전용 설비 → 휴머노이드 기반 유연 생산 전환",
            "중국 업체 가격 경쟁력 감안한 조기 도입 전략 수립",
        ],
        "color": LG_RED,
    },
    {
        "title": "🧠 AI 스마트팩토리",
        "subtitle": "공정 내 AI 인라인 검사 확산",
        "obs": [
            "GPU 엣지 컴퓨팅 기반 실시간 불량 인라인 검출",
            "디지털 트윈으로 공정 파라미터 AI 최적화 플랫폼",
            "RaaS(Robot as a Service) 구독 모델 확산",
            "하드웨어 대비 SW 플랫폼 가치 이동 뚜렷",
        ],
        "tips": [
            "자율제조(Autonomous Manufacturing) 로드맵 수립",
            "AI 비전 검사 솔루션 공급업체 비교 평가 착수",
            "외주 솔루션 vs. 자체 AI 플랫폼 구축 최적점 도출",
            "보유 공정 데이터 → AI 모델 학습 파이프라인 정비",
        ],
        "color": LG_NAVY,
    },
    {
        "title": "⚙ 7축 로봇 용접",
        "subtitle": "특수 제조 공정 고도화",
        "obs": [
            "6축 대비 높은 자유도 — 복잡한 곡면 용접 가능",
            "좁은 공간 접근성 우수, 장애물 회피 유연성",
            "가전 외장재·전장 부품 정밀 용접 적용 가능성",
            "용접 품질 AI 실시간 모니터링 통합 솔루션 등장",
        ],
        "tips": [
            "대형 가전(세탁기·냉장고) 금속 프레임 용접 공정 적용 검토",
            "기존 설비 인터페이스(OPC-UA, PROFINET) 호환성 검증",
            "용접 파라미터 AI 최적화 결합 시 품질 향상 기대",
        ],
        "color": RGBColor(0x1A, 0x7A, 0x3C),
    },
]

for i, c in enumerate(cards):
    x = Inches(0.25 + i * 4.36)
    add_rect(sl, x, Inches(1.15), Inches(4.1), Inches(5.85), fill_rgb=LG_LGRAY)
    tag_box(sl, x, Inches(1.15), Inches(4.1), Inches(0.5),
            c["title"], fill=c["color"])
    add_text(sl, c["subtitle"], x+Inches(0.1), Inches(1.7), Inches(3.9), Inches(0.35),
             font_size=11, bold=True, color=LG_GRAY)
    add_text(sl, "▮ 관찰 내용", x+Inches(0.1), Inches(2.1), Inches(3.9), Inches(0.3),
             font_size=12, bold=True, color=LG_NAVY)
    add_multiline(sl, c["obs"], x+Inches(0.12), Inches(2.4), Inches(3.86), Inches(1.65),
                  font_size=11, color=LG_DARK)
    add_rect(sl, x+Inches(0.1), Inches(4.1), Inches(3.9), Inches(0.03), fill_rgb=c["color"])
    add_text(sl, "▮ 시사점", x+Inches(0.1), Inches(4.15), Inches(3.9), Inches(0.3),
             font_size=12, bold=True, color=c["color"])
    add_multiline(sl, c["tips"], x+Inches(0.12), Inches(4.47), Inches(3.86), Inches(2.1),
                  font_size=11, color=LG_DARK)


# ═══════════════════════════════════════════════════
# SLIDE 6 — 산업 트렌드 분석
# ═══════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill_rgb=WHITE)
slide_header(sl, "산업 트렌드 분석",
             "이번 전시회에서 확인한 3대 메가 트렌드")

trends = [
    {
        "num": "01",
        "title": "로봇 · AI · 모빌리티\n산업 융합 가속",
        "body": [
            "자동차·모빌리티 산업과 로봇·자동화 산업의",
            "경계가 허물어지는 트렌드 명확히 확인",
            "",
            "EV 배터리 조립 로봇, 전장 부품 제조 자동화,",
            "자율주행 기술의 제조 현장 적용이 하나의",
            "생태계로 수렴 중",
            "",
            "→ 부산모빌리티쇼 동시 개최가 단순 행사",
            "   규모 확장 아닌 산업 융합 반영",
        ],
        "color": LG_RED,
    },
    {
        "num": "02",
        "title": "중국 로봇 기업\n기술 고도화 충격",
        "body": [
            "유니트리(Unitree), DJI 등 중국 기업의",
            "기술 수준이 글로벌 선두권에 도달",
            "",
            "가격 경쟁력은 국내 대비 압도적 우위",
            "",
            "국내 기업 대응 방향:",
            "  • 차별화된 SW 플랫폼",
            "  • 도메인 특화 솔루션",
            "  • 고신뢰성 산업(방산·항공·반도체)",
            "    집중 전략 불가피",
        ],
        "color": LG_NAVY,
    },
    {
        "num": "03",
        "title": "소프트웨어 · AI 중심\n가치 이동",
        "body": [
            "하드웨어(로봇 본체) 대비 Fleet Mgmt SW,",
            "AI 제어 플랫폼, 클라우드 연동 서비스의",
            "중요성이 부각",
            "",
            "로봇 구매 → RaaS(Robot as a Service)",
            "구독 모델 확산 추세 확인",
            "",
            "→ '어떤 로봇을 사느냐'보다",
            "  '어떤 플랫폼으로 운용하느냐'가",
            "  경쟁 우위 결정 요소로 부상",
        ],
        "color": RGBColor(0x7B, 0x3F, 0x9A),
    },
]

for i, t in enumerate(trends):
    x = Inches(0.25 + i * 4.36)
    add_rect(sl, x, Inches(1.15), Inches(4.1), Inches(5.85), fill_rgb=LG_LGRAY)
    # 번호 박스
    add_rect(sl, x, Inches(1.15), Inches(4.1), Inches(0.75), fill_rgb=t["color"])
    add_text(sl, t["num"], x+Inches(0.12), Inches(1.18), Inches(0.8), Inches(0.7),
             font_size=30, bold=True, color=WHITE)
    add_text(sl, t["title"], x+Inches(0.92), Inches(1.22), Inches(3.1), Inches(0.65),
             font_size=14, bold=True, color=WHITE)
    add_multiline(sl, t["body"], x+Inches(0.15), Inches(2.0), Inches(3.82), Inches(4.5),
                  font_size=12, color=LG_DARK)


# ═══════════════════════════════════════════════════
# SLIDE 7 — 시사점 및 제언
# ═══════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill_rgb=WHITE)
slide_header(sl, "시사점 및 제언", "LG전자 생산기술원 단기 실행 과제 & 중장기 전략")

# ── 상단: 단기 실행 과제 ──
tag_box(sl, Inches(0.25), Inches(1.15), Inches(12.8), Inches(0.45),
        "단기 실행 과제 (6개월 이내)", fill=LG_RED)

short_term = [
    ("1순위", "협동로봇 파일럿 공정 적용 검토 및 PoC 기획", "공정자동화팀", LG_RED),
    ("2순위", "AMR 도입 대상 공장(창원·구미) 선정 및 타당성 분석", "물류혁신팀", LG_NAVY),
    ("3순위", "AI 비전 검사 솔루션 공급업체 비교 평가", "품질기술팀", RGBColor(0x1A, 0x7A, 0x3C)),
    ("4순위", "휴머노이드 로봇 기술 로드맵 내부 검토 착수", "미래기술팀", RGBColor(0x7B, 0x3F, 0x9A)),
]
for i, (rank, task, dept, color) in enumerate(short_term):
    y = Inches(1.65 + i * 0.55)
    add_rect(sl, Inches(0.25), y, Inches(0.95), Inches(0.48), fill_rgb=color)
    add_text(sl, rank, Inches(0.25), y+Inches(0.06), Inches(0.95), Inches(0.38),
             font_size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(sl, Inches(1.22), y, Inches(9.4), Inches(0.48), fill_rgb=LG_LGRAY)
    add_text(sl, task, Inches(1.3), y+Inches(0.08), Inches(9.2), Inches(0.38),
             font_size=12, color=LG_DARK)
    add_rect(sl, Inches(10.64), y, Inches(2.4), Inches(0.48), fill_rgb=color)
    add_text(sl, dept, Inches(10.64), y+Inches(0.06), Inches(2.4), Inches(0.38),
             font_size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ── 하단: 중장기 전략 ──
add_rect(sl, Inches(0.25), Inches(4.0), Inches(12.8), Inches(0.04), fill_rgb=LG_RED)
tag_box(sl, Inches(0.25), Inches(4.08), Inches(12.8), Inches(0.45),
        "중장기 전략 방향 (1~3년)", fill=LG_NAVY)

mid_items = [
    ("① 자율제조 로드맵", "2028년 목표 공장 자동화율 설정 및 단계별 추진 계획 수립"),
    ("② 휴머노이드 파일럿", "물류·보조 작업부터 단계적 도입, 2027년 파일럿 라인 구성 목표"),
    ("③ 계열사 연계 강화", "로보스타 AGV/AMR 기술 연계 — 외부 도입 vs. 내부 개발 최적점 도출"),
    ("④ 기술 모니터링 체계", "분기별 국내외 주요 전시회 관람 및 기술 동향 내재화 체계 구축"),
]
for i, (title, desc) in enumerate(mid_items):
    x = Inches(0.25 + (i % 2) * 6.42)
    y = Inches(4.6 + (i // 2) * 0.95)
    add_rect(sl, x, y, Inches(6.2), Inches(0.82), fill_rgb=LG_LGRAY)
    add_rect(sl, x, y, Inches(0.06), Inches(0.82), fill_rgb=LG_RED)
    add_text(sl, title, x+Inches(0.15), y+Inches(0.04), Inches(5.9), Inches(0.32),
             font_size=12, bold=True, color=LG_NAVY)
    add_text(sl, desc, x+Inches(0.15), y+Inches(0.38), Inches(5.9), Inches(0.4),
             font_size=11, color=LG_DARK)


# ═══════════════════════════════════════════════════
# SLIDE 8 — 결론
# ═══════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill_rgb=LG_NAVY)
add_rect(sl, 0, 0, Inches(0.3), H, fill_rgb=LG_RED)
add_rect(sl, 0, Inches(5.5), W, Inches(2.0), fill_rgb=LG_RED)

add_text(sl, "결 론",
         Inches(0.6), Inches(0.8), Inches(10), Inches(0.5),
         font_size=18, color=LG_ACCENT)

add_multiline(sl, [
    "2026 부산로봇엑스포는 로봇·자동화 기술의 현 주소와 향후 방향성을",
    "파악하는 데 있어 매우 유의미한 전시회였습니다.",
], Inches(0.6), Inches(1.45), Inches(12.3), Inches(0.9),
   font_size=22, bold=True, color=WHITE)

add_text(sl, "3대 핵심 발견",
         Inches(0.6), Inches(2.5), Inches(5), Inches(0.4),
         font_size=15, bold=True, color=LG_ACCENT)

for i, txt in enumerate([
    "휴머노이드 로봇의 빠른 상용화 진행 — 예상보다 이른 제조 현장 투입 시점",
    "AMR 기반 물류 자동화의 성숙 — 즉시 도입 가능한 기술 완성도 확인",
    "AI 품질검사의 인라인 적용 확산 — SW 플랫폼 중심 경쟁 구도 형성",
]):
    add_rect(sl, Inches(0.6), Inches(3.0 + i * 0.5), Inches(0.35), Inches(0.38),
             fill_rgb=WHITE)
    add_text(sl, txt,
             Inches(1.1), Inches(3.02 + i * 0.5), Inches(11.5), Inches(0.38),
             font_size=14, color=WHITE)

add_text(sl,
         "LG전자 생산기술원 2026  |  2026 부산로봇기술산업전시회(ROBO EXPO Busan 2026)",
         Inches(0.6), Inches(5.7), Inches(11), Inches(0.4),
         font_size=13, bold=True, color=LG_NAVY)


# ─────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────
out = "/home/user/segamario-fluid-dynamics/roboexpo_2026_관람보고서_LG생산기술원.pptx"
prs.save(out)
print(f"Saved: {out}")
