"""
Builds the Toddler Potty Training Tracker as a real .xlsx workbook (openpyxl),
same pattern as build_planner_xlsx.py in pregnancy-to-toddler-planner/ so it
uploads straight to Google Drive and auto-converts into a native Google Sheet.
Checkbox columns land as TRUE/FALSE booleans — in Google Sheets, select each
"Done"/"Achieved" column and use Insert > Checkbox to make them interactive.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter

PASTEL = {
    'pink': 'FBD8E5', 'pinkDark': 'F4AFC9',
    'blue': 'D6E8F7', 'blueDark': 'AFD3EF',
    'mint': 'D9F2E6', 'mintDark': 'AEE4C9',
    'lavender': 'E6DFF7', 'lavenderDark': 'C9B8ED',
    'yellow': 'FFF3CE', 'yellowDark': 'FBE192',
    'peach': 'FDE3D3', 'peachDark': 'F8C39E',
    'gray': 'F4F4F4', 'textDark': '4A4A4A',
    'warnBg': 'F9D6D6', 'warnText': 'B33A3A',
    'okBg': 'DFF3E3', 'okText': '2F7A4D',
}


def fill(hexcode):
    return PatternFill(start_color="FF" + hexcode, end_color="FF" + hexcode, fill_type="solid")


def font(size=11, bold=False, italic=False, color=PASTEL['textDark']):
    return Font(size=size, bold=bold, italic=italic, color="FF" + color)


def banner(ws, ncols, title, subtitle, bg, row=1):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=title)
    c.fill = fill(bg); c.font = font(20, True); c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 32
    if subtitle:
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=ncols)
        c2 = ws.cell(row=row + 1, column=1, value=subtitle)
        c2.fill = fill(bg); c2.font = font(12, italic=True); c2.alignment = Alignment(horizontal='center')
        ws.row_dimensions[row + 1].height = 20


def header_row(ws, row, headers, bg=PASTEL['gray']):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(11, True); c.fill = fill(bg)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def note_row(ws, row, ncols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(10, italic=True); c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 34


def start_date_row(ws, row, label, ncols, note_text):
    ws.cell(row=row, column=1, value=label).font = font(11, True)
    dc = ws.cell(row=row, column=2, value="2026-01-01")
    dc.number_format = 'yyyy-mm-dd'
    dc.fill = fill(PASTEL['yellow'])
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=3, value=note_text)
    c.font = font(10, italic=True)


def checklist_sheet(wb, name, tab_color, banner_bg, title, subtitle, sections, note=None, date_col=False):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab_color
    ncols = 4 if date_col else 3
    headers = ['Item', 'Done', 'Date', 'Notes'] if date_col else ['Item / Task', 'Done', 'Notes']
    banner(ws, ncols, title, subtitle, banner_bg)
    header_row(ws, 3, headers)
    r = 4
    for sec in sections:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(row=r, column=1, value=sec['label'])
        c.font = font(11, True); c.fill = fill(sec.get('color', banner_bg))
        r += 1
        for item in sec['items']:
            ws.cell(row=r, column=1, value=item)
            ws.cell(row=r, column=2, value=False)
            if date_col:
                dc = ws.cell(row=r, column=3, value=None)
                dc.number_format = 'yyyy-mm-dd'
            r += 1
    if note:
        note_row(ws, r + 1, ncols, note)
    ws.column_dimensions['A'].width = 46
    ws.column_dimensions['B'].width = 10
    if date_col:
        ws.column_dimensions['C'].width = 16
        ws.column_dimensions['D'].width = 34
    else:
        ws.column_dimensions['C'].width = 38
    return ws


wb = openpyxl.Workbook()

# ---------- 1. Welcome ----------
ws = wb.active
ws.title = 'Welcome'
ws.sheet_properties.tabColor = PASTEL['lavenderDark']
banner(ws, 3, '🚽 Toddler Potty Training Tracker', 'A calm, encouraging companion for ages 2–4', PASTEL['lavender'])
rows = [
    ('How to use each tab', '', True),
    ('Readiness Checklist', 'The physical, behavioral and skill signs worth watching for before you start.', False),
    ('Daily Potty Log', 'One row per day — successes, accidents, and how it went.', False),
    ('Reward Chart', 'A sticker-style chart plus a reward menu to redeem stickers against.', False),
    ('Nighttime Training', 'A separate, later-stage log — nighttime dryness usually comes months after daytime.', False),
    ('Regression Log', 'For the setbacks that come with new siblings, moves, illness, or big changes.', False),
    ('Accident Patterns', 'Log time and place — the chart below helps you spot patterns to plan around.', False),
    ('Milestones', 'The bigger checkpoints, celebrated in whatever order they happen.', False),
    ('Dashboard', 'A one-page snapshot of progress so far.', False),
    ('', '', False),
    ('Getting started', '', True),
    ('1. Skim the Readiness Checklist — there is no rush, and no single "right" age to start.', '', False),
    ('2. Go to the Daily Potty Log tab and enter your start date at the top.', '', False),
    ('3. Log daily. Duplicate the sheet (right-click tab > Duplicate) if you need more rows.', '', False),
    ('4. Select each "Done"/"Achieved" column and use Insert > Checkbox to make them clickable.', '', False),
    ('5. Regressions are normal, not failure — the Regression Log tab is there for exactly that.', '', False),
]
r = 3
for label, desc, bold in rows:
    c = ws.cell(row=r, column=1, value=label)
    if bold:
        c.font = font(13, True)
    ws.cell(row=r, column=2, value=desc)
    r += 1
note_row(ws, r + 1, 3, '💡 Every child is ready at a different age, and progress is rarely a straight line. This tracker is a place to keep notes, not a schedule to keep up with.')
ws.column_dimensions['A'].width = 42
ws.column_dimensions['B'].width = 60
ws.column_dimensions['C'].width = 20

# ---------- 2. Readiness Checklist ----------
sections = [
    {'label': 'Physical Signs', 'color': PASTEL['pink'], 'items': [
        'Stays dry for 2+ hours at a time', 'Has regular, somewhat predictable bowel movements',
        'Can pull pants/underwear up and down', 'Shows discomfort with a wet or dirty diaper',
        'Can walk to and sit on the potty independently',
    ]},
    {'label': 'Behavioral Signs', 'color': PASTEL['blue'], 'items': [
        'Shows interest in the bathroom or toilet', 'Wants to wear "big kid" underwear',
        'Tells you before or during an accident', 'Hides or goes still when pooping',
        'Can sit still for 2–5 minutes', 'Wants more independence generally',
    ]},
    {'label': 'Skill Signs', 'color': PASTEL['mint'], 'items': [
        'Can follow simple two-step instructions', 'Can communicate needs (words, signs, or gestures)',
        'Shows interest in imitating adults/older siblings',
    ]},
]
checklist_sheet(wb, 'Readiness', PASTEL['pinkDark'], PASTEL['pink'],
                '✅ Readiness Checklist', 'Signs worth watching for — not a deadline', sections,
                "Tip: most kids show some but not all of these before starting, and that's fine — this is a general reference, not a rule.")

# ---------- 3. Daily Potty Log ----------
ws = wb.create_sheet('Daily Potty Log')
ws.sheet_properties.tabColor = PASTEL['blueDark']
headers = ['Day', 'Date', 'Successes (#)', 'Accidents (#)', 'Self-Initiated (#)', 'Dry Through Nap', 'Reward Given', 'Notes']
banner(ws, len(headers), '📝 Daily Potty Log', 'One row per day', PASTEL['blue'])
start_date_row(ws, 3, 'Start date:', len(headers), '⬅ Enter once — the Date column below fills in automatically.')
header_row(ws, 4, headers)
first = 5
n = 30
for i in range(n):
    day = i + 1
    row = first + i
    ws.cell(row=row, column=1, value=day)
    dc = ws.cell(row=row, column=2, value=f"=$B$3+({day}-1)")
    dc.number_format = 'yyyy-mm-dd'
    ws.cell(row=row, column=7, value=False)
dv_yn = DataValidation(type="list", formula1='"Y,N,N/A"', allow_blank=True)
ws.add_data_validation(dv_yn)
dv_yn.add(f"F{first}:F{first + n - 1}")
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 16
ws.column_dimensions['H'].width = 32
note_row(ws, first + n + 1, len(headers),
         'Tip: count successes and accidents, not perfection — the trend over 2–3 weeks tells you more than any single day.')

# ---------- 4. Reward Chart ----------
ws = wb.create_sheet('Reward Chart')
ws.sheet_properties.tabColor = PASTEL['yellowDark']
headers = ['Day', 'Date', 'Sticker Earned', 'Reward Redeemed', 'Notes']
banner(ws, len(headers), '⭐ Reward Chart', 'A sticker for a good try, not just a perfect one', PASTEL['yellow'])
start_date_row(ws, 3, 'Start date:', len(headers), '⬅ Same start date as the Daily Potty Log tab.')
header_row(ws, 4, headers)
first = 5
n = 30
for i in range(n):
    day = i + 1
    row = first + i
    ws.cell(row=row, column=1, value=day)
    dc = ws.cell(row=row, column=2, value=f"=$B$3+({day}-1)")
    dc.number_format = 'yyyy-mm-dd'
    ws.cell(row=row, column=3, value=False)
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 18
ws.column_dimensions['E'].width = 30

reward_start = first + n + 2
ws.merge_cells(start_row=reward_start, start_column=1, end_row=reward_start, end_column=len(headers))
c = ws.cell(row=reward_start, column=1, value='Reward Menu — fill in with your child')
c.font = font(12, True); c.fill = fill(PASTEL['peach'])
header_row(ws, reward_start + 1, ['Stickers Needed', 'Reward', '', '', ''], bg=PASTEL['gray'])
reward_rows = [(5, ''), (10, ''), (20, '')]
for i, (n_stickers, reward) in enumerate(reward_rows):
    row = reward_start + 2 + i
    ws.cell(row=row, column=1, value=n_stickers)
    ws.cell(row=row, column=2, value=reward)
note_row(ws, reward_start + 2 + len(reward_rows) + 1, len(headers),
         'Tip: small, frequent rewards (a sticker, a high-five, a favorite song) tend to work better than big, rare ones at this age.')

# ---------- 5. Nighttime Training ----------
ws = wb.create_sheet('Nighttime Training')
ws.sheet_properties.tabColor = PASTEL['lavenderDark']
headers = ['Day', 'Date', 'Dry Overnight', 'Woke to Use Potty (#)', 'Wearing', 'Notes']
banner(ws, len(headers), '🌙 Nighttime Training', 'Usually starts months after daytime training — no rush', PASTEL['lavender'])
start_date_row(ws, 3, 'Start date:', len(headers), '⬅ This is often a separate, later start date from daytime training.')
header_row(ws, 4, headers)
first = 5
n = 30
for i in range(n):
    day = i + 1
    row = first + i
    ws.cell(row=row, column=1, value=day)
    dc = ws.cell(row=row, column=2, value=f"=$B$3+({day}-1)")
    dc.number_format = 'yyyy-mm-dd'
dv_dry = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
dv_wearing = DataValidation(type="list", formula1='"Underwear,Pull-Up,Diaper"', allow_blank=True)
ws.add_data_validation(dv_dry); dv_dry.add(f"C{first}:C{first + n - 1}")
ws.add_data_validation(dv_wearing); dv_wearing.add(f"E{first}:E{first + n - 1}")
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 18
ws.column_dimensions['F'].width = 30
note_row(ws, first + n + 1, len(headers),
         "Tip: many kids stay in pull-ups overnight well after daytime training — that's typical, not a setback.")

# ---------- 6. Regression Log ----------
ws = wb.create_sheet('Regression Log')
ws.sheet_properties.tabColor = PASTEL['peachDark']
headers = ['Date', 'Possible Trigger', 'What Happened', 'How We Responded', 'Notes']
banner(ws, len(headers), '🔄 Regression Log', 'Setbacks happen — this is a place to notice why, not to worry', PASTEL['peach'])
header_row(ws, 3, headers)
first = 4
n = 15
for i in range(n):
    row = first + i
    dc = ws.cell(row=row, column=1, value=None)
    dc.number_format = 'yyyy-mm-dd'
dv_trigger = DataValidation(type="list",
                             formula1='"New sibling,Move/travel,Illness,New school/daycare,Big schedule change,Stress,Unknown,Other"',
                             allow_blank=True)
ws.add_data_validation(dv_trigger)
dv_trigger.add(f"B{first}:B{first + n - 1}")
for col, w in zip(range(1, 6), [16, 22, 30, 30, 30]):
    ws.column_dimensions[get_column_letter(col)].width = w
note_row(ws, first + n + 1, len(headers),
         'Regressions are extremely common and almost always temporary. A calm, low-pressure response tends to help more than correction.')

# ---------- 7. Accident Patterns ----------
ws = wb.create_sheet('Accident Patterns')
ws.sheet_properties.tabColor = PASTEL['mintDark']
headers = ['Date', 'Time of Day', 'Location', 'Likely Cause', 'Notes']
banner(ws, len(headers), '📍 Accident Patterns', 'Time and place tell you more than the accident itself', PASTEL['mint'])
header_row(ws, 3, headers)
first = 4
n = 30
for i in range(n):
    row = first + i
    dc = ws.cell(row=row, column=1, value=None)
    dc.number_format = 'yyyy-mm-dd'
dv_time = DataValidation(type="list", formula1='"Morning,Afternoon,Evening,Overnight"', allow_blank=True)
dv_loc = DataValidation(type="list", formula1='"Home,Car,Preschool/Daycare,Out & About,Other"', allow_blank=True)
ws.add_data_validation(dv_time); dv_time.add(f"B{first}:B{first + n - 1}")
ws.add_data_validation(dv_loc); dv_loc.add(f"C{first}:C{first + n - 1}")
for col, w in zip(range(1, 6), [16, 16, 20, 22, 30]):
    ws.column_dimensions[get_column_letter(col)].width = w

summary_row = first + n + 2
ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=5)
c = ws.cell(row=summary_row, column=1, value='Accidents by Time of Day (auto-counted from above)')
c.font = font(12, True); c.fill = fill(PASTEL['gray'])
times = ['Morning', 'Afternoon', 'Evening', 'Overnight']
for i, t in enumerate(times):
    row = summary_row + 1 + i
    ws.cell(row=row, column=1, value=t)
    ws.cell(row=row, column=2, value=f'=COUNTIF(B{first}:B{first + n - 1},A{row})')
chart = BarChart()
chart.title = 'Accidents by Time of Day'
data_ref = Reference(ws, min_col=2, min_row=summary_row + 1, max_row=summary_row + len(times))
cats_ref = Reference(ws, min_col=1, min_row=summary_row + 1, max_row=summary_row + len(times))
chart.add_data(data_ref, titles_from_data=False)
chart.set_categories(cats_ref)
ws.add_chart(chart, f"D{summary_row + 1}")
note_row(ws, summary_row + len(times) + 12, len(headers),
         'Tip: a cluster around one time or place (e.g. right after arriving at daycare, or during car rides) usually means a proactive reminder there helps more than a reaction after the fact.')

# ---------- 8. Milestones ----------
ws = wb.create_sheet('Milestones')
ws.sheet_properties.tabColor = PASTEL['lavenderDark']
banner(ws, 4, '🎉 Milestones', 'In whatever order they happen', PASTEL['lavender'])
header_row(ws, 3, ['Milestone', 'Achieved', 'Date Achieved', 'Notes'])
milestones = [
    'Shows interest in the potty/toilet', 'Sits on the potty with clothes on',
    'Sits on the potty with clothes off', 'First pee in the potty', 'First poop in the potty',
    'Tells us before an accident', 'Initiates going to the potty independently',
    'Stays dry through a nap', 'Wipes with help', 'Wipes independently',
    'Washes hands independently', 'Fully out of daytime diapers/pull-ups',
    'Stays dry overnight', 'Fully out of nighttime diapers/pull-ups',
]
r = 4
for m in milestones:
    ws.cell(row=r, column=1, value=m)
    ws.cell(row=r, column=2, value=False)
    dc = ws.cell(row=r, column=3, value=None)
    dc.number_format = 'yyyy-mm-dd'
    r += 1
ws.column_dimensions['A'].width = 42
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 16
ws.column_dimensions['D'].width = 34
note_row(ws, r + 1, 4, 'Daytime and nighttime dryness often arrive months apart — both are worth celebrating on their own.')

# ---------- 9. Dashboard ----------
ws = wb.create_sheet('Dashboard')
ws.sheet_properties.tabColor = PASTEL['pinkDark']
banner(ws, 8, '📊 Dashboard', 'A one-page snapshot of progress so far', PASTEL['peach'])

daily_ws = wb['Daily Potty Log']
accident_ws = wb['Accident Patterns']

chart1 = LineChart()
chart1.title = 'Successes vs. Accidents by Day'
s_ref = Reference(daily_ws, min_col=3, min_row=5, max_row=34)
a_ref = Reference(daily_ws, min_col=4, min_row=5, max_row=34)
cats1 = Reference(daily_ws, min_col=1, min_row=5, max_row=34)
chart1.add_data(s_ref, titles_from_data=False)
chart1.add_data(a_ref, titles_from_data=False)
chart1.set_categories(cats1)
chart1.series[0].tx = openpyxl.chart.series.SeriesLabel(v="Successes")
chart1.series[1].tx = openpyxl.chart.series.SeriesLabel(v="Accidents")
ws.add_chart(chart1, "A3")

chart2 = BarChart()
chart2.title = 'Accidents by Time of Day'
acc_summary_first = 4 + 30 + 3
data_ref2 = Reference(accident_ws, min_col=2, min_row=acc_summary_first, max_row=acc_summary_first + 3)
cats_ref2 = Reference(accident_ws, min_col=1, min_row=acc_summary_first, max_row=acc_summary_first + 3)
chart2.add_data(data_ref2, titles_from_data=False)
chart2.set_categories(cats_ref2)
ws.add_chart(chart2, "H3")

ws.cell(row=22, column=1, value='Milestones achieved:').font = font(11, True)
ws.cell(row=23, column=1, value=f'=COUNTIF(Milestones!B:B,TRUE)&" of {len(milestones)} checked"').font = font(18, True, color=PASTEL['pinkDark'])
ws.cell(row=25, column=1, value='Stickers earned so far:').font = font(11, True)
ws.cell(row=26, column=1, value="=COUNTIF('Reward Chart'!C:C,TRUE)").font = font(18, True, color=PASTEL['pinkDark'])
for col in range(1, 9):
    ws.column_dimensions[get_column_letter(col)].width = 16

wb.save('Toddler_Potty_Training_Tracker.xlsx')
print('wrote Toddler_Potty_Training_Tracker.xlsx')
