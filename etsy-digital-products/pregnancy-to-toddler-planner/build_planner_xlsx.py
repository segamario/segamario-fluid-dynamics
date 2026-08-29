"""
Builds the Bump to Toddler Planner as a real .xlsx workbook (openpyxl),
mirroring build_planner.gs, so it can be uploaded straight to Google Drive
and auto-converted into a native Google Sheet (no Apps Script execution
needed). Checkbox columns land as TRUE/FALSE booleans — in Google Sheets,
select each "Done"/"Achieved"/"Given" column and use Insert > Checkbox to
turn them into real interactive checkboxes (one-time, ~10 columns total).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter

PASTEL = {
    'pink': 'FBD8E5', 'pinkDark': 'F4AFC9',
    'blue': 'D6E8F7', 'blueDark': 'AFD3EF',
    'mint': 'D9F2E6', 'mintDark': 'AEE4C9',
    'lavender': 'E6DFF7', 'lavenderDark': 'C9B8ED',
    'yellow': 'FFF3CE', 'yellowDark': 'FBE192',
    'peach': 'FDE3D3', 'peachDark': 'F8C39E',
    'gray': 'F4F4F4', 'textDark': '4A4A4A',
    'warnBg': 'F9D6D6', 'warnText': 'B33A3A',
    'okBg': 'DFF3E3', 'okText': '2F7A4D',
}


def fill(hexcode):
    return PatternFill(start_color="FF" + hexcode, end_color="FF" + hexcode, fill_type="solid")


def font(size=11, bold=False, italic=False, color=PASTEL['textDark']):
    return Font(size=size, bold=bold, italic=italic, color="FF" + color)


def banner(ws, ncols, title, subtitle, bg, row=1):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=title)
    c.fill = fill(bg); c.font = font(20, True); c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 32
    if subtitle:
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=ncols)
        c2 = ws.cell(row=row + 1, column=1, value=subtitle)
        c2.fill = fill(bg); c2.font = font(12, italic=True); c2.alignment = Alignment(horizontal='center')
        ws.row_dimensions[row + 1].height = 20


def header_row(ws, row, headers, bg=PASTEL['gray']):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(11, True); c.fill = fill(bg)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def note_row(ws, row, ncols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(10, italic=True); c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 34


def due_date_row(ws, row, label, ncols, note_text):
    ws.cell(row=row, column=1, value=label).font = font(11, True)
    dc = ws.cell(row=row, column=2, value="2026-01-01")
    dc.number_format = 'yyyy-mm-dd'
    dc.fill = fill(PASTEL['yellow'])
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=3, value=note_text)
    c.font = font(10, italic=True)


def checklist_sheet(wb, name, tab_color, banner_bg, title, subtitle, sections, note=None):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab_color
    banner(ws, 3, title, subtitle, banner_bg)
    header_row(ws, 3, ['Item / Task', 'Done', 'Notes'])
    r = 4
    for sec in sections:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c = ws.cell(row=r, column=1, value=sec['label'])
        c.font = font(11, True); c.fill = fill(sec.get('color', banner_bg))
        r += 1
        for item in sec['items']:
            ws.cell(row=r, column=1, value=item)
            ws.cell(row=r, column=2, value=False)
            r += 1
    if note:
        note_row(ws, r + 1, 3, note)
    ws.column_dimensions['A'].width = 46
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 38
    return ws


wb = openpyxl.Workbook()

# ---------- 1. Welcome ----------
ws = wb.active
ws.title = 'Welcome'
ws.sheet_properties.tabColor = PASTEL['lavenderDark']
banner(ws, 3, '🤱 Bump to Toddler Planner', 'Your complete pregnancy-to-age-2 care companion', PASTEL['lavender'])
rows = [
    ('How to use each tab', '', True),
    ('Pregnancy Prep Checklist', 'Trimester-by-trimester to-dos, from first appointment to hospital pre-registration.', False),
    ('Prenatal Health & Nutrition Tracker', 'Log weekly weight, blood pressure, vitamins and symptoms. Enter your due date once — dates fill in automatically.', False),
    ('Hospital Bag Checklist', 'Everything to pack for mom, baby, and your support person.', False),
    ('Postpartum Mom Recovery Tracker', 'Daily recovery log for the first 6 weeks after birth.', False),
    ('Newborn Daily Care Log (0-3mo)', 'Feeding, diaper, sleep and temperature tracking for the newborn stage.', False),
    ('Baby Weekly Log (3-24mo)', 'A lighter-touch weekly log once routines settle in.', False),
    ('Baby Growth Tracker', 'Weight / height / head circumference at each well-check visit, plotted automatically.', False),
    ('Developmental Milestones', 'General milestone reference by age band, 0-24 months.', False),
    ('Vaccination Schedule', 'A general reference checklist — always confirm the real schedule with your pediatrician.', False),
    ('Dashboard', 'A one-page snapshot of the charts above.', False),
    ('', '', False),
    ('Getting started', '', True),
    ('1. Go to the Prenatal Health tab and enter your due date at the top.', '', False),
    ('2. Go to the Newborn Daily Log / Baby Weekly Log tabs and enter baby’s birth date at the top once baby arrives.', '', False),
    ('3. Check off items as you go.', '', False),
    ('4. Duplicate the Newborn Daily Log or Baby Weekly Log sheet (right-click tab > Duplicate) if you need more rows.', '', False),
    ('5. Select each "Done"/"Achieved"/"Given" column and use Insert > Checkbox to make them clickable.', '', False),
]
r = 3
for label, desc, bold in rows:
    c = ws.cell(row=r, column=1, value=label)
    if bold:
        c.font = font(13, True)
    ws.cell(row=r, column=2, value=desc)
    r += 1
note_row(ws, r + 1, 3, '💡 This planner is for organizational purposes only and does not replace professional medical advice. Always consult your OB/GYN or pediatrician for anything health-related.')
ws.column_dimensions['A'].width = 42
ws.column_dimensions['B'].width = 60
ws.column_dimensions['C'].width = 20

# ---------- 2. Pregnancy Prep Checklist ----------
sections = [
    {'label': 'First Trimester (Weeks 1–13)', 'color': PASTEL['pink'], 'items': [
        'Confirm pregnancy & schedule first prenatal visit', 'Choose an OB/GYN or midwife',
        'Start prenatal vitamins (folic acid, iron, DHA)', 'Share family medical history with your provider',
        'Cut out alcohol, smoking, and unpasteurized/raw foods',
        'Track early symptoms (nausea, fatigue) in the Prenatal Tracker tab',
        'Look into maternity leave / workplace policies',
        'Tell your employer (per your comfort level & policy deadlines)'
    ]},
    {'label': 'Second Trimester (Weeks 14–27)', 'color': PASTEL['blue'], 'items': [
        'Schedule anatomy scan / mid-pregnancy ultrasound', 'Complete glucose screening (gestational diabetes)',
        'Research childbirth / breastfeeding classes', 'Start nursery planning',
        'Create a gift registry / plan a baby shower', 'Buy maternity clothes', 'Choose a pediatrician',
        'Consider cord blood banking (optional)', 'Start a prenatal exercise / pelvic floor routine',
        'Draft a baby name shortlist'
    ]},
    {'label': 'Third Trimester (Weeks 28–40)', 'color': PASTEL['mint'], 'items': [
        'Pack hospital bag (see Hospital Bag Checklist tab)', 'Install car seat and get it inspected',
        'Write and share your birth plan', 'Pre-register at the hospital',
        'Finish nursery & run a crib/room safety check', 'Wash baby clothes & linens', 'Prep and freeze meals',
        'Arrange help for the first weeks postpartum', 'Pack documents: ID, insurance card, birth plan copies',
        'Confirm postpartum support (doula, family, partner leave)', 'Finalize maternity leave start date'
    ]},
]
checklist_sheet(wb, 'Pregnancy Prep', PASTEL['pinkDark'], PASTEL['pink'],
                '🤰 Pregnancy Prep Checklist', 'Trimester by trimester', sections,
                'Tip: aim to have the hospital bag and car seat done by week 36 — babies don’t always wait for your due date.')

# ---------- 3. Prenatal Health & Nutrition Tracker ----------
ws = wb.create_sheet('Prenatal Health')
ws.sheet_properties.tabColor = PASTEL['mintDark']
headers = ['Week', 'Date', 'Weight', 'Blood Pressure', 'Swelling', 'Felt Baby Move', 'Prenatal Vitamin',
           'Water (cups)', 'Iron-Rich Food', 'Calcium-Rich Food', 'Symptoms / Notes']
banner(ws, len(headers), '🥗 Prenatal Health & Nutrition Tracker', 'Weekly log, weeks 1–40', PASTEL['mint'])
due_date_row(ws, 3, 'Due date:', len(headers), '⬅ Enter your due date once — the Date column below fills in automatically.')
header_row(ws, 4, headers)
first = 5
n = 40
for i in range(n):
    week = i + 1
    row = first + i
    ws.cell(row=row, column=1, value=week)
    dc = ws.cell(row=row, column=2, value=f"=$B$3-((40-{week})*7)")
    dc.number_format = 'yyyy-mm-dd'
    ws.cell(row=row, column=7, value=False)
    ws.cell(row=row, column=9, value=False)
    ws.cell(row=row, column=10, value=False)
dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
ws.add_data_validation(dv_yn)
dv_yn.add(f"E{first}:F{first + n - 1}")
for col in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11):
    ws.column_dimensions[get_column_letter(col)].width = 16
ws.column_dimensions['K'].width = 32
note_row(ws, first + n + 1, len(headers),
         '⚠️ Call your provider right away for: heavy bleeding, severe headache, sudden swelling, reduced fetal movement, or fever.')

# ---------- 4. Hospital Bag Checklist ----------
sections = [
    {'label': 'For Mom', 'color': PASTEL['pink'], 'items': [
        'Comfortable going-home outfit', 'Nursing bra & nursing pads', 'Maternity/postpartum pads',
        'Toiletries & travel-size essentials', 'Phone charger (long cord)', 'Slippers & robe',
        'Snacks & drinks', 'Nipple cream', 'Going-home outfit (loose, comfy)'
    ]},
    {'label': 'For Baby', 'color': PASTEL['blue'], 'items': [
        'Going-home outfits (2 sizes, newborn + 0-3mo)', 'Swaddle blankets', 'Newborn diapers',
        'Car seat installed & inspected', 'Hat & mittens', 'Burp cloths'
    ]},
    {'label': 'For Partner / Support Person', 'color': PASTEL['mint'], 'items': [
        'Change of clothes', 'Snacks', 'Entertainment (book, headphones)', 'Camera / charger', 'Pillow'
    ]},
    {'label': 'Documents & Extras', 'color': PASTEL['yellow'], 'items': [
        'Photo ID', 'Insurance card', 'Hospital pre-registration forms', 'Printed birth plan copies',
        'Pediatrician contact info', 'Car seat manual (for hospital staff check)'
    ]},
]
checklist_sheet(wb, 'Hospital Bag', PASTEL['blueDark'], PASTEL['blue'],
                '🎒 Hospital Bag & Newborn Essentials', 'Pack by week 36', sections,
                'Tip: keep this bag by the door from week 36 on — check it off as you pack, not as you leave.')

# ---------- 5. Postpartum Mom Recovery Tracker ----------
ws = wb.create_sheet('Postpartum Mom')
ws.sheet_properties.tabColor = PASTEL['pinkDark']
headers = ['Day', 'Date', 'Bleeding / Lochia', 'Pain Level (1-5)', 'Temperature (°F)', 'Mood', 'Sleep (hrs)',
           'Feeding Method', 'Rest Taken', 'Notes']
banner(ws, len(headers), '💗 Postpartum Mom Recovery Tracker', 'Daily log, first 6 weeks', PASTEL['pink'])
due_date_row(ws, 3, 'Birth date:', len(headers), '⬅ Enter baby’s birth date once — dates below fill in automatically.')
header_row(ws, 4, headers)
first = 5
n = 42
for i in range(n):
    day = i + 1
    row = first + i
    ws.cell(row=row, column=1, value=day)
    dc = ws.cell(row=row, column=2, value=f"=$B$3+({day}-1)")
    dc.number_format = 'yyyy-mm-dd'
    ws.cell(row=row, column=9, value=False)
dv_bleed = DataValidation(type="list", formula1='"Heavy,Moderate,Light,Spotting,None"', allow_blank=True)
dv_pain = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
dv_mood = DataValidation(type="list", formula1='"Great,OK,Low,Struggling"', allow_blank=True)
dv_feed = DataValidation(type="list", formula1='"Breastfeeding,Bottle,Mixed"', allow_blank=True)
for dv, col in ((dv_bleed, 'C'), (dv_pain, 'D'), (dv_mood, 'F'), (dv_feed, 'H')):
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{first + n - 1}")
ws.conditional_formatting.add(
    f"E{first}:E{first + n - 1}",
    CellIsRule(operator='greaterThanOrEqual', formula=['100.4'], fill=fill(PASTEL['warnBg']), font=font(11, color=PASTEL['warnText']))
)
ws.conditional_formatting.add(
    f"F{first}:F{first + n - 1}",
    FormulaRule(formula=[f'F{first}="Struggling"'], fill=fill(PASTEL['warnBg']), font=font(11, color=PASTEL['warnText']))
)
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 15
ws.column_dimensions['J'].width = 32
note_row(ws, first + n + 1, len(headers),
         '⚠️ Contact your provider right away for: fever ≥100.4°F, heavy bleeding, severe pain, or persistent low mood — you deserve support, not just survival.')

# ---------- 6. Newborn Daily Care Log ----------
ws = wb.create_sheet('Newborn Daily Log')
ws.sheet_properties.tabColor = PASTEL['blueDark']
headers = ['Day', 'Date', 'Feedings (#)', 'Feeding Method', 'Wet Diapers (#)', 'Dirty Diapers (#)', 'Stool Color',
           'Sleep (hrs)', 'Temp (°F)', 'Notes']
banner(ws, len(headers), '🍼 Newborn Daily Care Log', 'Birth to ~3 months — one row per day', PASTEL['blue'])
due_date_row(ws, 3, 'Baby’s birth date:', len(headers), '⬅ Enter once — duplicate this tab for month 2 / month 3.')
header_row(ws, 4, headers)
first = 5
n = 30
for i in range(n):
    day = i + 1
    row = first + i
    ws.cell(row=row, column=1, value=day)
    dc = ws.cell(row=row, column=2, value=f"=$B$3+({day}-1)")
    dc.number_format = 'yyyy-mm-dd'
dv_feedmethod = DataValidation(type="list", formula1='"Breast,Bottle,Both"', allow_blank=True)
dv_stool = DataValidation(type="list", formula1='"Yellow/Seedy,Green,Brown,Black (first days),Red-streaked,N/A"', allow_blank=True)
ws.add_data_validation(dv_feedmethod); dv_feedmethod.add(f"D{first}:D{first + n - 1}")
ws.add_data_validation(dv_stool); dv_stool.add(f"G{first}:G{first + n - 1}")
ws.conditional_formatting.add(
    f"E{first}:E{first + n - 1}",
    CellIsRule(operator='lessThan', formula=['6'], fill=fill(PASTEL['peachDark']), font=font(11, color=PASTEL['warnText']))
)
ws.conditional_formatting.add(
    f"I{first}:I{first + n - 1}",
    CellIsRule(operator='greaterThanOrEqual', formula=['100.4'], fill=fill(PASTEL['warnBg']), font=font(11, color=PASTEL['warnText']))
)
ws.conditional_formatting.add(
    f"G{first}:G{first + n - 1}",
    FormulaRule(formula=[f'G{first}="Red-streaked"'], fill=fill(PASTEL['warnBg']), font=font(11, color=PASTEL['warnText']))
)
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 15
ws.column_dimensions['J'].width = 30
note_row(ws, first + n + 1, len(headers),
         '⚠️ Call your pediatrician for: fewer than 6 wet diapers/day after day 5, fever ≥100.4°F (call immediately under 3 months), or black/red stool after the first days.')

# ---------- 7. Baby Weekly Log ----------
ws = wb.create_sheet('Baby Weekly Log')
ws.sheet_properties.tabColor = PASTEL['mintDark']
headers = ['Week #', 'Age (months)', 'Week Of', 'Avg Feedings/Day', 'Avg Sleep (hrs/day)', 'Diapers/Day (approx)',
           'Naps/Day', 'New This Week', 'Notes']
banner(ws, len(headers), '📅 Baby Weekly Log', '3–24 months — a lighter check-in once routines settle', PASTEL['mint'])
due_date_row(ws, 3, 'Baby’s birth date:', len(headers), '⬅ Same birth date as the Newborn Daily Log tab.')
header_row(ws, 4, headers)
first = 5
start_week, end_week = 13, 103
for i, wk in enumerate(range(start_week, end_week + 1)):
    row = first + i
    ws.cell(row=row, column=1, value=wk)
    ws.cell(row=row, column=2, value=f"=ROUND({wk}/4.345,1)")
    dc = ws.cell(row=row, column=3, value=f"=$B$3+(({wk}-1)*7)")
    dc.number_format = 'yyyy-mm-dd'
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 16
ws.column_dimensions['I'].width = 30
note_row(ws, first + (end_week - start_week + 1) + 1, len(headers),
         'Tip: fill this in weekly instead of daily once feeding and sleep settle into a routine, usually around 3 months.')

# ---------- 8. Baby Growth Tracker ----------
ws = wb.create_sheet('Growth Tracker')
ws.sheet_properties.tabColor = PASTEL['lavenderDark']
headers = ['Check-up', 'Age', 'Date', 'Weight', 'Height / Length', 'Head Circumference']
banner(ws, len(headers), '📏 Baby Growth Tracker', 'Weight, height & head circumference at each well-check', PASTEL['lavender'])
header_row(ws, 3, headers)
checkups = ['Birth', '2 Weeks', '1 Month', '2 Months', '4 Months', '6 Months', '9 Months', '12 Months',
            '15 Months', '18 Months', '21 Months', '24 Months']
first = 4
for i, c in enumerate(checkups):
    row = first + i
    ws.cell(row=row, column=1, value=c)
    ws.cell(row=row, column=2, value=c)
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 20

chart = LineChart()
chart.title = 'Growth Curve — Weight & Height'
last = first + len(checkups) - 1
weight_ref = Reference(ws, min_col=4, min_row=first, max_row=last)
height_ref = Reference(ws, min_col=5, min_row=first, max_row=last)
cats = Reference(ws, min_col=1, min_row=first, max_row=last)
chart.add_data(weight_ref, titles_from_data=False)
chart.add_data(height_ref, titles_from_data=False)
chart.set_categories(cats)
chart.series[0].tx = openpyxl.chart.series.SeriesLabel(v="Weight")
chart.series[1].tx = openpyxl.chart.series.SeriesLabel(v="Height")
ws.add_chart(chart, f"A{last + 3}")
note_row(ws, last + 22, len(headers),
         'Enter the numbers from your pediatrician visits here — the chart above updates automatically.')

# ---------- 9. Developmental Milestones ----------
ws = wb.create_sheet('Milestones')
ws.sheet_properties.tabColor = PASTEL['yellowDark']
banner(ws, 4, '🧠 Developmental Milestones', 'General reference, 0–24 months', PASTEL['yellow'])
header_row(ws, 3, ['Milestone', 'Achieved', 'Date Achieved', 'Notes'])
bands = [
    {'label': '0–3 Months', 'color': PASTEL['pink'], 'items': [
        'Lifts head briefly during tummy time', 'Follows faces/objects with eyes', 'Smiles responsively',
        'Makes cooing sounds', 'Brings hands to mouth', 'Startles at loud sounds'
    ]},
    {'label': '4–6 Months', 'color': PASTEL['blue'], 'items': [
        'Rolls from tummy to back', 'Holds head steady without support', 'Reaches for and grabs objects',
        'Laughs out loud', 'Begins babbling ("ba-ba", "da-da")', 'Sits with support'
    ]},
    {'label': '7–9 Months', 'color': PASTEL['mint'], 'items': [
        'Sits without support', 'Crawls or scoots', 'Transfers objects hand to hand',
        'Responds to own name', 'Plays peek-a-boo', 'Picks up small objects (pincer grasp)'
    ]},
    {'label': '10–12 Months', 'color': PASTEL['lavender'], 'items': [
        'Pulls to stand', 'Cruises along furniture', 'May take first independent steps',
        'Says 1–2 words with meaning', 'Waves "bye-bye"', 'Points at objects of interest'
    ]},
    {'label': '13–18 Months', 'color': PASTEL['peach'], 'items': [
        'Walks independently', 'Says several single words', 'Follows simple one-step instructions',
        'Stacks 2–3 blocks', 'Drinks from an open cup', 'Points to a body part when asked'
    ]},
    {'label': '19–24 Months', 'color': PASTEL['pink'], 'items': [
        'Runs', 'Kicks a ball', 'Uses 2-word phrases', 'Follows simple 2-step instructions',
        'Shows growing independence', 'Begins pretend play'
    ]},
]
r = 4
for band in bands:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    c = ws.cell(row=r, column=1, value=band['label'])
    c.font = font(11, True); c.fill = fill(band['color'])
    r += 1
    for item in band['items']:
        ws.cell(row=r, column=1, value=item)
        ws.cell(row=r, column=2, value=False)
        dc = ws.cell(row=r, column=3, value=None)
        dc.number_format = 'yyyy-mm-dd'
        r += 1
ws.column_dimensions['A'].width = 46
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 16
ws.column_dimensions['D'].width = 34
note_row(ws, r + 1, 4, 'Every baby develops at their own pace. This is a general reference, not a diagnostic tool — talk to your pediatrician with any concerns.')

# ---------- 10. Vaccination Schedule ----------
ws = wb.create_sheet('Vaccinations')
ws.sheet_properties.tabColor = PASTEL['lavenderDark']
banner(ws, 6, '💉 Vaccination Schedule Tracker', 'General reference — confirm with your pediatrician', PASTEL['lavender'])
header_row(ws, 3, ['Vaccine', 'Recommended Age', 'Dose #', 'Given', 'Date Given', 'Notes'])
vax_rows = [
    ['Hepatitis B', 'Birth', '1'], ['Hepatitis B', '1–2 Months', '2'],
    ['DTaP', '2 Months', '1'], ['IPV (Polio)', '2 Months', '1'], ['Hib', '2 Months', '1'],
    ['PCV13 (Pneumococcal)', '2 Months', '1'], ['Rotavirus', '2 Months', '1'],
    ['DTaP', '4 Months', '2'], ['IPV (Polio)', '4 Months', '2'], ['Hib', '4 Months', '2'],
    ['PCV13 (Pneumococcal)', '4 Months', '2'], ['Rotavirus', '4 Months', '2'],
    ['DTaP', '6 Months', '3'], ['Hib', '6 Months', '3'], ['PCV13 (Pneumococcal)', '6 Months', '3'],
    ['Rotavirus', '6 Months', '3'], ['Hepatitis B', '6–18 Months', '3'],
    ['Influenza', '6+ Months', 'Annual'], ['MMR', '12–15 Months', '1'],
    ['Varicella (Chickenpox)', '12–15 Months', '1'], ['Hepatitis A', '12–23 Months', '1'],
    ['PCV13 (Pneumococcal)', '12–15 Months', 'Booster'], ['Hib', '12–15 Months', 'Booster'],
    ['DTaP', '15–18 Months', 'Booster'], ['Hepatitis A', '18–23 Months', '2'],
]
first = 4
for i, row_vals in enumerate(vax_rows):
    row = first + i
    ws.cell(row=row, column=1, value=row_vals[0])
    ws.cell(row=row, column=2, value=row_vals[1])
    ws.cell(row=row, column=3, value=row_vals[2])
    ws.cell(row=row, column=4, value=False)
    dc = ws.cell(row=row, column=5, value=None)
    dc.number_format = 'yyyy-mm-dd'
for col, w in zip(range(1, 7), [26, 18, 12, 10, 16, 30]):
    ws.column_dimensions[get_column_letter(col)].width = w
note_row(ws, first + len(vax_rows) + 1, 6,
         '⚠️ Vaccination schedules vary by country and change over time. This is a general reference only — always follow your pediatrician’s or local health authority’s official schedule.')

# ---------- 11. Dashboard ----------
ws = wb.create_sheet('Dashboard')
ws.sheet_properties.tabColor = PASTEL['pinkDark']
banner(ws, 8, '📊 Dashboard', 'A one-page snapshot of your journey', PASTEL['peach'])

prenatal_ws = wb['Prenatal Health']
growth_ws = wb['Growth Tracker']
newborn_ws = wb['Newborn Daily Log']

chart1 = LineChart()
chart1.title = 'Maternal Weight by Week'
w_ref = Reference(prenatal_ws, min_col=3, min_row=5, max_row=44)
cats1 = Reference(prenatal_ws, min_col=1, min_row=5, max_row=44)
chart1.add_data(w_ref, titles_from_data=False)
chart1.set_categories(cats1)
ws.add_chart(chart1, "A3")

chart2 = LineChart()
chart2.title = 'Baby Growth Curve'
w2 = Reference(growth_ws, min_col=4, min_row=4, max_row=15)
h2 = Reference(growth_ws, min_col=5, min_row=4, max_row=15)
cats2 = Reference(growth_ws, min_col=1, min_row=4, max_row=15)
chart2.add_data(w2, titles_from_data=False)
chart2.add_data(h2, titles_from_data=False)
chart2.set_categories(cats2)
ws.add_chart(chart2, "H3")

chart3 = BarChart()
chart3.title = 'Newborn Feedings & Wet Diapers (Month 1)'
feed_ref = Reference(newborn_ws, min_col=3, min_row=5, max_row=34)
wet_ref = Reference(newborn_ws, min_col=5, min_row=5, max_row=34)
cats3 = Reference(newborn_ws, min_col=1, min_row=5, max_row=34)
chart3.add_data(feed_ref, titles_from_data=False)
chart3.add_data(wet_ref, titles_from_data=False)
chart3.set_categories(cats3)
ws.add_chart(chart3, "A22")

ws.cell(row=22, column=8, value='Milestones achieved:').font = font(11, True)
ws.cell(row=23, column=8, value='=COUNTIF(Milestones!B:B,TRUE)&" of 36 checked"').font = font(18, True, color=PASTEL['pinkDark'])
for col in range(1, 9):
    ws.column_dimensions[get_column_letter(col)].width = 16

wb.save('Bump_to_Toddler_Planner.xlsx')
print('wrote Bump_to_Toddler_Planner.xlsx')
